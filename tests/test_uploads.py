import hashlib
import io
import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import Mock, patch
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from docx import Document as WordDocument
from PIL import Image
import requests
from rest_framework.test import APITestCase

from ai_core.models import Document, DocumentChunk
from ai_core.services.cloudinary_service import cloudinary_asset_from_url, get_upload_signature, is_allowed_remote_url
from ai_core.services.docx_parser import DocxNativeParser
from exams.models import ImageBank, Question, QuestionImage
from tests.support import seed_lms


ASSET_URL = 'https://res.cloudinary.com/test/raw/upload/v123/tests/bài.docx'


def docx_bytes():
    doc = WordDocument()
    doc.add_paragraph('Câu 1. Một cộng một bằng bao nhiêu?')
    doc.add_paragraph('A. 2').runs[0].bold = True
    doc.add_paragraph('B. 3')
    doc.add_table(rows=1, cols=2).cell(0, 0).text = 'Bảng đáp án'
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class CloudinaryBoundaryTests(SimpleTestCase):
    @patch('ai_core.services.cloudinary_service.socket.getaddrinfo')
    def test_only_https_cloudinary_with_public_dns_is_allowed(self, dns):
        dns.return_value = [(2, 1, 6, '', ('104.16.1.1', 443))]
        self.assertTrue(is_allowed_remote_url(ASSET_URL))
        for url in ('http://res.cloudinary.com/x', 'https://cloudinary.com.evil.test/x',
                    'https://127.0.0.1/x', 'file:///etc/passwd', 'https://example.com/x'):
            with self.subTest(url=url):
                self.assertFalse(is_allowed_remote_url(url))
        for address in ('127.0.0.1', '10.0.0.1', '169.254.169.254', '::1', 'fc00::1'):
            dns.return_value = [(2, 1, 6, '', (address, 443))]
            with self.subTest(address=address):
                self.assertFalse(is_allowed_remote_url(ASSET_URL))
        dns.side_effect = OSError('DNS unavailable')
        self.assertFalse(is_allowed_remote_url(ASSET_URL))

    def test_public_ids_preserve_folders_and_raw_extensions(self):
        for url, expected in (
            (ASSET_URL, ('tests/bài.docx', 'raw')),
            ('https://res.cloudinary.com/test/image/upload/v123/tests/a%20b.png', ('tests/a b', 'image')),
            ('https://res.cloudinary.com/test/raw/upload/tests/a.pdf', ('tests/a.pdf', 'raw')),
            ('https://example.com', (None, None)),
        ):
            with self.subTest(url=url):
                self.assertEqual(cloudinary_asset_from_url(url), expected)

    @patch.dict(os.environ, {'CLOUDINARY_API_KEY': 'test-key', 'CLOUDINARY_API_SECRET': 'test-secret'})
    @patch('ai_core.services.cloudinary_service.time.time', return_value=123)
    def test_signature_matches_contract_without_leaking_secret(self, _):
        signature = get_upload_signature()
        self.assertEqual(signature['signature'], hashlib.sha1(b'timestamp=123test-secret').hexdigest())
        self.assertEqual(signature['api_key'], 'test-key')
        self.assertNotIn('test-secret', str(signature))

    @patch.dict(os.environ, {'CLOUDINARY_API_KEY': '', 'CLOUDINARY_API_SECRET': ''})
    def test_missing_credentials_do_not_produce_signature(self):
        self.assertIsNone(get_upload_signature())


class DocxParserTests(SimpleTestCase):
    def test_real_docx_text_formatting_and_table(self):
        with tempfile.NamedTemporaryFile(suffix='.docx') as file:
            file.write(docx_bytes())
            file.flush()
            blocks = DocxNativeParser.parse_docx(file.name)
        self.assertIn('Một cộng một', str(blocks))
        self.assertIn('Bảng đáp án', str(blocks))
        self.assertTrue(any(block.get('fmt', {}).get('bold') for block in blocks if 'A. 2' in str(block)))

    def test_corrupt_empty_oversized_and_xml_entities_rejected(self):
        for content in (b'', b'not a zip'):
            with tempfile.NamedTemporaryFile(suffix='.docx') as file:
                file.write(content)
                file.flush()
                with self.assertRaises(zipfile.BadZipFile):
                    DocxNativeParser.parse_docx(file.name)
        with patch('ai_core.services.docx_parser.os.path.getsize', return_value=25 * 1024 * 1024 + 1):
            with self.assertRaises(ValueError):
                DocxNativeParser.parse_docx('unused.docx')
        with tempfile.NamedTemporaryFile(suffix='.docx') as file:
            with zipfile.ZipFile(file, 'w') as archive:
                archive.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY a "ENTITY">]><x>&a;</x>')
            file.flush()
            from defusedxml.common import DefusedXmlException
            with self.assertRaises(DefusedXmlException):
                DocxNativeParser.parse_docx(file.name)

    def test_docx_uncompressed_size_limit(self):
        with tempfile.NamedTemporaryFile(suffix='.docx') as file:
            file.write(docx_bytes())
            file.flush()
            with patch('ai_core.services.docx_parser.MAX_DOCX_UNCOMPRESSED_BYTES', 1):
                with self.assertRaises(ValueError):
                    DocxNativeParser.parse_docx(file.name)


class DocumentUploadTests(APITestCase):
    def setUp(self):
        self.d = seed_lms()
        self.client.force_authenticate(self.d.teacher)
        self.url = f'/api/ai/classes/{self.d.classroom.pk}/documents/upload/'
        self.temp = tempfile.TemporaryDirectory(prefix='nvh-upload-check-')
        self.addCleanup(self.temp.cleanup)
        original = tempfile.NamedTemporaryFile
        temp_patch = patch('ai_core.views.tempfile.NamedTemporaryFile',
                           side_effect=lambda **kwargs: original(dir=self.temp.name, **kwargs))
        temp_patch.start()
        self.addCleanup(temp_patch.stop)
        self.http_patch = patch('requests.get')
        self.http = self.http_patch.start()
        self.addCleanup(self.http_patch.stop)
        self.http.return_value.iter_content.return_value = [b'valid document bytes']
        dns = patch('ai_core.services.cloudinary_service.socket.getaddrinfo',
                    return_value=[(2, 1, 6, '', ('104.16.1.1', 443))])
        dns.start()
        self.addCleanup(dns.stop)
        deletion = patch('ai_core.views.delete_from_cloudinary')
        self.delete = deletion.start()
        self.addCleanup(deletion.stop)

    def upload(self, **extra):
        return self.client.post(self.url, {'file_url': ASSET_URL, 'file_name': 'bài.docx', **extra}, format='json')

    def assert_clean(self):
        self.assertEqual(list(Path(self.temp.name).iterdir()), [])

    @patch('ai_core.views.AIGeneratorService.ingest_document')
    def test_url_upload_persists_document_chunks_and_removes_temp_file(self, ingest):
        def process(path, doc_id):
            self.assertTrue(Path(path).exists())
            DocumentChunk.objects.create(document_id=doc_id, chunk_index=0, content='Nội dung lớp')
            return {'knowledge_chunks_count': 1, 'questions': []}
        ingest.side_effect = process
        res = self.upload()
        self.assertEqual(res.status_code, 201, res.data)
        doc = Document.objects.get()
        self.assertEqual(doc.cloudinary_public_id, 'tests/bài.docx')
        self.assertEqual(doc.chunks.count(), 1)
        self.assertEqual(res.data['knowledge_chunks_count'], 1)
        self.http.assert_called_once_with(ASSET_URL, stream=True, timeout=30, allow_redirects=False)
        self.delete.assert_not_called()
        self.assert_clean()

    @patch('ai_core.views.AIGeneratorService.ingest_document', return_value={'knowledge_chunks_count': 0, 'questions': []})
    @patch('ai_core.services.cloudinary_service.upload_to_cloudinary')
    def test_multipart_uses_raw_upload_and_removes_temp(self, upload, ingest):
        upload.return_value = {'secure_url': ASSET_URL, 'public_id': 'tests/bài.docx'}
        res = self.client.post(self.url, {'file': SimpleUploadedFile('bài.docx', docx_bytes())}, format='multipart')
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(upload.call_args.kwargs['resource_type'], 'raw')
        ingest.assert_called_once()
        self.assert_clean()

    @patch('ai_core.views.AIGeneratorService.ingest_document')
    def test_ingestion_failure_cleans_database_chunks_asset_and_temp(self, ingest):
        def fail(path, doc_id):
            DocumentChunk.objects.create(document_id=doc_id, chunk_index=0, content='partial')
            raise ValueError('invalid document')
        ingest.side_effect = fail
        self.assertEqual(self.upload().status_code, 500)
        self.assertFalse(Document.objects.exists())
        self.assertFalse(DocumentChunk.objects.exists())
        self.delete.assert_called_once_with('tests/bài.docx', resource_type='raw')
        self.assert_clean()

    def test_download_failure_does_not_leak_temp_file(self):
        self.http.side_effect = requests.Timeout('injected timeout')
        self.assertEqual(self.upload().status_code, 500)
        self.assertFalse(Document.objects.exists())
        self.assert_clean()

    @patch('ai_core.views.MAX_REMOTE_FILE_BYTES', 3)
    def test_stream_over_limit_does_not_leak_temp_file(self):
        self.http.return_value.iter_content.return_value = [b'123', b'4']
        self.assertEqual(self.upload().status_code, 500)
        self.assertFalse(Document.objects.exists())
        self.assert_clean()

    @patch('ai_core.views.AIGeneratorService.ingest_document')
    def test_invalid_url_missing_file_and_non_owner_fail_before_download(self, ingest):
        self.assertEqual(self.client.post(self.url, {}, format='json').status_code, 400)
        self.assertEqual(self.upload(file_url='https://127.0.0.1/a.docx').status_code, 400)
        for user in (self.d.other_teacher, self.d.student, self.d.outsider):
            self.client.force_authenticate(user)
            self.assertEqual(self.upload().status_code, 403)
        self.http.assert_not_called()
        ingest.assert_not_called()

    @patch('ai_core.views.AIGeneratorService.ingest_document', return_value={'questions': []})
    def test_size_limit_including_exact_boundary(self, ingest):
        with patch('ai_core.views.MAX_REMOTE_FILE_BYTES', 3):
            for content, expected in ((b'12', 201), (b'123', 201), (b'1234', 400)):
                with self.subTest(length=len(content)):
                    upload = patch('ai_core.services.cloudinary_service.upload_to_cloudinary',
                                   return_value={'secure_url': ASSET_URL, 'public_id': 'tests/bài.docx'})
                    with upload:
                        res = self.client.post(self.url, {'file': SimpleUploadedFile('bài.docx', content)}, format='multipart')
                    self.assertEqual(res.status_code, expected, res.data)
        self.assert_clean()

    @patch('ai_core.views.AIGeneratorService.ingest_document')
    def test_redirect_is_rejected_without_ingesting_body(self, ingest):
        ingest.return_value = {'questions': [], 'knowledge_chunks_count': 0}
        self.http.return_value.status_code = 302
        self.http.return_value.is_redirect = True
        self.http.return_value.headers = {'Location': 'http://127.0.0.1/private'}
        res = self.upload()
        self.assertGreaterEqual(res.status_code, 400)
        ingest.assert_not_called()
        self.assert_clean()


class ImportAndImageTests(APITestCase):
    def setUp(self):
        self.d = seed_lms()
        self.client.force_authenticate(self.d.teacher)

    @patch('ai_core.views.is_allowed_remote_url', return_value=True)
    @patch('requests.Session.get')
    @patch('ai_core.views.AIGeneratorService.extract_from_file')
    def test_batch_extract_deduplicates_questions_and_preserves_answers(self, extract, http, _):
        http.return_value.iter_content.return_value = [docx_bytes()]
        question = {'text': 'Nhận định', 'question_type': 'true_false',
                    'options': [{'text': 'Ý a', 'is_correct': False}]}
        extract.return_value = ([question], {})
        res = self.client.post('/api/ai/generate/extract-file/', {'files': [
            {'file_url': ASSET_URL, 'file_name': 'bài.docx'},
            {'file_url': ASSET_URL, 'file_name': 'bài.docx'},
        ]}, format='json')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['questions'], [question])
        self.assertEqual(extract.call_count, 2)
        self.assertFalse(Document.objects.exists())

    @patch('ai_core.views.is_allowed_remote_url', return_value=True)
    @patch('ai_core.views.delete_from_cloudinary')
    @patch('requests.Session.get')
    @patch('ai_core.views.AIGeneratorService.extract_from_file')
    def test_failed_batch_cleans_all_uploaded_assets(self, extract, http, delete, _):
        http.return_value.iter_content.return_value = [b'file']
        extract.side_effect = [([{'text': 'Câu đã đọc'}], {}), ValueError('bad second file')]
        res = self.client.post('/api/ai/generate/extract-file/', {'files': [
            {'file_url': ASSET_URL, 'file_name': 'bài.docx'},
            {'file_url': ASSET_URL.replace('bài.docx', 'second.pdf'), 'file_name': 'second.pdf'},
        ]}, format='json')
        self.assertEqual(res.status_code, 500)
        self.assertEqual(delete.call_count, 2)
        self.assertFalse(Question.objects.exclude(pk=self.d.mc.pk).exists())

    def test_batch_input_limits_and_types(self):
        for payload in ({}, {'files': ['bad']}, {'files': [{}] * 11}, {'file_url': 'file:///etc/passwd'}):
            with self.subTest(payload=payload):
                res = self.client.post('/api/ai/generate/extract-file/', payload, format='json')
                self.assertEqual(res.status_code, 400, res.data)

    def test_bulk_save_preserves_all_question_types_and_reports_rejections(self):
        drafts = [{'text': 'MC', 'options': [{'text': 'A', 'is_correct': True}]},
                  {'text': 'TF', 'question_type': 'true_false', 'context': 'Bối cảnh',
                   'options': [{'text': 'a', 'is_correct': False}]},
                  {'text': 'SA', 'question_type': 'short_answer', 'correct_answer_text': '42'},
                  {'text': ''}]
        res = self.client.post('/api/ai/generate/save-bulk/', {
            'subject_id': self.d.subject.pk, 'quiz_id': self.d.quiz.pk, 'questions': drafts,
        }, format='json')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(len(res.data['saved_ids']), 3)
        self.assertEqual(res.data['rejected'], [{'index': 3, 'reason': 'missing_question_content'}])
        self.assertEqual(self.d.quiz.quiz_questions.count(), 4)
        self.assertEqual(set(self.d.quiz.quiz_questions.values_list('points', flat=True)), {2.5})
        self.assertFalse(Question.objects.get(text='TF').options.get().is_correct)
        self.assertEqual(Question.objects.get(text='SA').correct_answer_text, '42')

    @patch('ai_core.views.Option.objects.create', side_effect=RuntimeError('injected option failure'))
    def test_bulk_save_failure_does_not_leave_orphan_question(self, _):
        before = Question.objects.count()
        res = self.client.post('/api/ai/generate/save-bulk/', {
            'subject_id': self.d.subject.pk, 'questions': [{'text': 'Partial', 'options': [{'text': 'a'}]}],
        }, format='json')
        self.assertEqual(res.status_code, 500)
        self.assertEqual(Question.objects.count(), before)

    def test_bulk_save_forbidden_on_another_teachers_quiz(self):
        self.client.force_authenticate(self.d.other_teacher)
        res = self.client.post('/api/ai/generate/save-bulk/', {
            'subject_id': self.d.subject.pk, 'quiz_id': self.d.quiz.pk, 'questions': [{'text': 'Unauthorized'}],
        }, format='json')
        self.assertEqual(res.status_code, 403)
        self.assertFalse(Question.objects.filter(text='Unauthorized').exists())

    def test_shared_image_survives_deleting_one_question(self):
        image = ImageBank.objects.create(sha256='a' * 64, image_file='https://res.cloudinary.com/test/image/upload/shared.png')
        other = Question.objects.create(text='Câu khác', subject=self.d.subject, created_by=self.d.teacher)
        for question in (self.d.mc, other):
            QuestionImage.objects.create(question=question, image=image, position=0)
        with patch('ai_core.services.cloudinary_service.delete_from_cloudinary') as delete:
            res = self.client.delete(f'/api/exams/questions/{self.d.mc.pk}/')
        self.assertEqual(res.status_code, 204)
        self.assertTrue(ImageBank.objects.filter(pk=image.pk).exists())
        self.assertTrue(other.question_images.exists())
        delete.assert_not_called()

    def test_invalid_image_and_student_upload_rejected(self):
        self.assertEqual(self.client.post('/api/exams/upload-image/', {
            'image': SimpleUploadedFile('fake.png', b'not an image', content_type='image/png'),
        }, format='multipart').status_code, 400)
        self.client.force_authenticate(self.d.student)
        self.assertEqual(self.client.get('/api/exams/cloudinary-signature/').status_code, 403)

    @patch('exams.serializers.QuestionSerializer._handle_images', side_effect=ValueError('injected failure'))
    def test_full_update_failure_keeps_original_question_and_options(self, _):
        with self.assertRaisesMessage(ValueError, 'injected failure'):
            self.client.put(f'/api/exams/questions/{self.d.mc.pk}/update-full/', {
                'subject': self.d.subject.pk, 'question_type': 'multiple_choice', 'text': 'Changed',
                'options': [{'text': 'Replacement', 'is_correct': True}], 'question_images': [],
            }, format='json')
        self.d.mc.refresh_from_db()
        self.assertEqual(self.d.mc.text, '1 + 1 bằng bao nhiêu?')
        self.assertEqual(self.d.mc.options.count(), 2)
